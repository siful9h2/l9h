#!/usr/bin/env python3
"""Build a formatted Excel copy of tools/update-log.csv: python3 tools/update_log_xlsx.py OUT.xlsx"""
import csv, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule

src = 'tools/update-log.csv'
out = sys.argv[1] if len(sys.argv) > 1 else 'L9H Website Update Log.xlsx'
rows = list(csv.reader(open(src, encoding='utf-8')))
head, data = rows[0], rows[1:]
F = 'Arial'
wb = Workbook(); ws = wb.active; ws.title = 'Update Log'
ws.append(head)
for r in data:
    r = list(r); r[0] = int(r[0])
    for k in (5, 6):
        r[k] = datetime.date.fromisoformat(r[k]) if r[k] else None
    ws.append(r)
n = len(data) + 1
widths = [6, 58, 18, 26, 26, 15, 13, 24, 52]
for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
for c in ws[1]:
    c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='3A2A18')
    c.alignment = Alignment(vertical='center', wrap_text=True)
ws.row_dimensions[1].height = 22
for row in ws.iter_rows(min_row=2, max_row=n):
    for c in row:
        c.font = Font(name=F, size=10); c.alignment = Alignment(vertical='top', wrap_text=True)
    row[5].number_format = row[6].number_format = 'mmm d, yyyy'
    row[0].alignment = Alignment(horizontal='center', vertical='top')
tab = Table(displayName='UpdateLog', ref=f'A1:I{n}')
tab.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
ws.add_table(tab); ws.freeze_panes = 'C2'
fills = {'Live': 'E3F1E1', 'Waiting for weekly publish': 'FFF1CC', 'Scheduled': 'E1ECF7', 'On hold': 'EEEEEE', 'Removed': 'F6E0E0', 'Proposal': 'F3E6F7'}
for status, color in fills.items():
    ws.conditional_formatting.add(f'H2:H{n}', FormulaRule(formula=[f'$H2="{status}"'], fill=PatternFill('solid', fgColor=color)))

s = wb.create_sheet('Summary')
s['A1'] = 'L9H website: update summary'; s['A1'].font = Font(name=F, bold=True, size=14)
s['A2'] = 'Counts come from the Update Log tab. The live log is tools/update-log.csv in the GitHub repo (siful9h2/l9h, staging branch).'
s['A2'].font = Font(name=F, size=9, italic=True, color='666666')
s['A4'], s['B4'] = 'Status', 'Updates'
for c in (s['A4'], s['B4']): c.font = Font(name=F, bold=True)
for i, st in enumerate(fills, start=5):
    s[f'A{i}'] = st; s[f'B{i}'] = f"=COUNTIF('Update Log'!$H$2:$H${n},A{i})"
    s[f'A{i}'].font = s[f'B{i}'].font = Font(name=F)
t = 5 + len(fills)
s[f'A{t}'] = 'Total'; s[f'B{t}'] = f'=SUM(B5:B{t-1})'
s[f'A{t}'].font = s[f'B{t}'].font = Font(name=F, bold=True)
s[f'A{t+2}'] = 'Most recent update developed'; s[f'B{t+2}'] = f"=MAX('Update Log'!$F$2:$F${n})"
s[f'A{t+3}'] = 'Most recent update live'; s[f'B{t+3}'] = f"=MAX('Update Log'!$G$2:$G${n})"
for r in (t + 2, t + 3):
    s[f'A{r}'].font = s[f'B{r}'].font = Font(name=F); s[f'B{r}'].number_format = 'mmm d, yyyy'
s.column_dimensions['A'].width = 32; s.column_dimensions['B'].width = 16
wb.save(out); print('wrote', out)
